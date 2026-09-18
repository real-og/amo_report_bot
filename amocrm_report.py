from __future__ import annotations

import argparse
import json
import os
import re
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


WON_STATUS_ID = 142
LOST_STATUS_ID = 143
HEADERS = [
    "Название сделки", "Этап сделки", "Контакт сделки",
    "Компания", "Ответственный менеджер", "Город", "Услуга", "Сумма",
    "Источник", "Примечания", "Причина нереализации",
]
MONTHS_RU = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
]


@dataclass(frozen=True)
class Stage:
    id: int
    name: str
    sort: int
    color: str


class AmoCRMClient:
    def __init__(self, env_path: Path):
        load_dotenv(env_path)
        subdomain = require_env("AMOCRM_SUBDOMAIN")
        self.base_url = f"https://{subdomain}.amocrm.ru"
        self.access_token = require_env("AMOCRM_ACCESS_TOKEN")
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": "AmoCRMDealsReport/1.0"})

    def request(self, method: str, path: str, **kwargs: Any) -> Any:
        url = f"{self.base_url}{path}"
        for attempt in range(5):
            headers = kwargs.pop("headers", {})
            headers["Authorization"] = f"Bearer {self.access_token}"
            response = self.session.request(method, url, headers=headers, timeout=60, **kwargs)
            if response.status_code == 401:
                raise RuntimeError("amoCRM отклонила долгосрочный токен: проверьте AMOCRM_ACCESS_TOKEN")
            if response.status_code == 429:
                time.sleep(min(2 ** attempt, 20))
                continue
            response.raise_for_status()
            return response.json() if response.content else None
        raise RuntimeError(f"Не удалось выполнить запрос {method} {path}")

    def get_all(self, path: str, entity: str, params: dict[str, Any] | None = None) -> list[dict]:
        result: list[dict] = []
        page = 1
        while True:
            query = dict(params or {})
            query.update({"page": page, "limit": 250})
            data = self.request("GET", path, params=query)
            batch = data.get("_embedded", {}).get(entity, []) if data else []
            result.extend(batch)
            if len(batch) < 250:
                break
            page += 1
        return result

    def pipelines(self) -> list[dict]:
        return self.get_all("/api/v4/leads/pipelines", "pipelines")

    def users(self) -> dict[int, str]:
        users = self.get_all("/api/v4/users", "users")
        return {item["id"]: item.get("name", "") for item in users}

    def leads(self, pipeline_id: int, date_from: int, date_to: int | None) -> list[dict]:
        params: dict[str, Any] = {
            "with": "contacts,loss_reason",
            "filter[pipeline_id][]": pipeline_id,
            "filter[created_at][from]": date_from,
        }
        if date_to is not None:
            params["filter[created_at][to]"] = date_to
        return self.get_all("/api/v4/leads", "leads", params)

    def contacts(self, ids: set[int]) -> dict[int, str]:
        return self._entity_names("/api/v4/contacts", "contacts", ids)

    def companies(self, ids: set[int]) -> dict[int, str]:
        return self._entity_names("/api/v4/companies", "companies", ids)

    def _entity_names(self, path: str, entity: str, ids: set[int]) -> dict[int, str]:
        result: dict[int, str] = {}
        for id_chunk in chunks(sorted(ids), 50):
            items = self.get_all(path, entity, {"filter[id][]": id_chunk})
            result.update({item["id"]: item.get("name", "") for item in items})
        return result

    def lead_notes(self, lead_ids: set[int]) -> dict[int, list[str]]:
        result: dict[int, list[tuple[int, str]]] = {}
        for id_chunk in chunks(sorted(lead_ids), 50):
            notes = self.get_all(
                "/api/v4/leads/notes",
                "notes",
                {"filter[entity_id][]": id_chunk},
            )
            for note in notes:
                params = note.get("params") or {}
                text = str(params.get("text") or "").strip()
                if note.get("note_type") != "common" or not text:
                    continue
                result.setdefault(note["entity_id"], []).append(
                    (note.get("created_at", 0), text)
                )
        return {
            lead_id: [text for _, text in sorted(items)]
            for lead_id, items in result.items()
        }


def require_env(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        raise RuntimeError(f"Не заполнена переменная {name} в .env")
    return value


def parse_date(value: str | None, end_of_day: bool = False) -> int | None:
    if not value:
        return None
    dt = datetime.strptime(value, "%Y-%m-%d")
    if end_of_day:
        dt = dt.replace(hour=23, minute=59, second=59)
    return int(dt.timestamp())


def chunks(items: list[int], size: int) -> Iterable[list[int]]:
    for start in range(0, len(items), size):
        yield items[start:start + size]


def custom_value(lead: dict, field_cfg: dict[str, Any]) -> str:
    target_id = field_cfg.get("id")
    target_name = str(field_cfg.get("name", "")).strip().casefold()
    for field in lead.get("custom_fields_values") or []:
        by_id = target_id is not None and field.get("field_id") == target_id
        by_name = target_name and str(field.get("field_name", "")).strip().casefold() == target_name
        if by_id or by_name:
            values = []
            for item in field.get("values") or []:
                value = item.get("value", "")
                if isinstance(value, dict):
                    value = value.get("name") or value.get("value") or json.dumps(value, ensure_ascii=False)
                if value not in (None, ""):
                    values.append(str(value))
            return ", ".join(values)
    return ""


def related_ids(lead: dict, entity: str) -> list[int]:
    return [
        item["id"]
        for item in lead.get("_embedded", {}).get(entity, [])
        if item.get("id") is not None
    ]


def related_names(lead: dict, entity: str, names: dict[int, str]) -> str:
    return ", ".join(
        filter(None, (names.get(item_id, "") for item_id in related_ids(lead, entity)))
    )


def tag_names(lead: dict) -> str:
    return ", ".join(
        item.get("name", "")
        for item in lead.get("_embedded", {}).get("tags", [])
        if item.get("name")
    )


def loss_reason_name(lead: dict) -> str:
    reason = lead.get("_embedded", {}).get("loss_reason")
    if isinstance(reason, list):
        return ", ".join(item.get("name", "") for item in reason if item.get("name"))
    if isinstance(reason, dict):
        return str(reason.get("name") or "")
    return ""


def sanitize_filename(value: str) -> str:
    return re.sub(r'[<>:"/\\|?*]+', "_", value).strip(" .") or "pipeline"


def pipeline_stages(pipeline: dict, config: dict) -> dict[int, Stage]:
    excluded = {x.casefold() for x in config.get("excluded_status_names", [])}
    included_names = [x.strip().casefold() for x in config.get("included_status_names", [])]
    configured_order = {name: index for index, name in enumerate(included_names)}
    palette = config["colors"]["default_stages"]
    stages: dict[int, Stage] = {}
    for status in sorted(pipeline.get("_embedded", {}).get("statuses", []), key=lambda x: x.get("sort", 0)):
        status_id = status["id"]
        name = status.get("name", str(status_id))
        normalized_name = name.strip().casefold()
        if normalized_name in excluded:
            continue
        if status_id == WON_STATUS_ID:
            color, order = config["colors"]["won"], 1_000_000
        elif status_id == LOST_STATUS_ID:
            color, order = config["colors"]["lost"], 1_000_001
        else:
            if included_names and normalized_name not in configured_order:
                continue
            order = configured_order.get(normalized_name, status.get("sort", 0))
            color = palette[order % len(palette)]
        stages[status_id] = Stage(status_id, name, order, color)

    actual_regular_names = {
        stage.name.strip().casefold()
        for stage in stages.values()
        if stage.id not in (WON_STATUS_ID, LOST_STATUS_ID)
    }
    missing = [name for name in included_names if name not in actual_regular_names]
    if missing:
        print(
            f"Предупреждение для воронки «{pipeline.get('name', pipeline.get('id'))}»: "
            f"не найдены этапы: {', '.join(missing)}"
        )
    return stages


def lead_row(
    lead: dict,
    stage: Stage,
    users: dict[int, str],
    fields: dict,
    contacts: dict[int, str],
    companies: dict[int, str],
    notes: dict[int, list[str]],
) -> list[Any]:
    return [
        lead.get("name", ""),
        stage.name,
        related_names(lead, "contacts", contacts),
        related_names(lead, "companies", companies),
        users.get(lead.get("responsible_user_id"), ""),
        custom_value(lead, fields["city"]),
        custom_value(lead, fields["service"]),
        lead.get("price") or 0,
        tag_names(lead),
        "\n\n".join(notes.get(lead["id"], [])),
        loss_reason_name(lead),
    ]


def month_key(lead: dict) -> tuple[int, int]:
    dt = datetime.fromtimestamp(lead["created_at"])
    return dt.year, dt.month


def create_workbook(
    pipeline_name: str,
    leads: list[dict],
    stages: dict[int, Stage],
    users: dict[int, str],
    contacts: dict[int, str],
    companies: dict[int, str],
    notes: dict[int, list[str]],
    config: dict,
    output: Path,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сделки"
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title = ws.cell(1, 1, f"Отчёт по воронке «{pipeline_name}»")
    title.font = Font(size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1F4E78")
    title.alignment = Alignment(horizontal="center")

    for column, header in enumerate(HEADERS, 1):
        cell = ws.cell(2, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}2"

    valid = [lead for lead in leads if lead.get("status_id") in stages]
    valid.sort(key=lambda lead: (month_key(lead), stages[lead["status_id"]].sort, lead.get("created_at", 0)))
    month_totals: dict[tuple[int, int], dict[str, float]] = {}
    for lead in valid:
        totals = month_totals.setdefault(month_key(lead), {"work": 0, "won": 0})
        amount = float(lead.get("price") or 0)
        if lead["status_id"] == WON_STATUS_ID:
            totals["won"] += amount
        elif lead["status_id"] != LOST_STATUS_ID:
            totals["work"] += amount

    current_month: tuple[int, int] | None = None
    row_number = 3
    thin = Side(style="thin", color="D9E2F3")

    for index, lead in enumerate(valid):
        key = month_key(lead)
        if key != current_month:
            current_month = key
            ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=len(HEADERS))
            month_cell = ws.cell(row_number, 1, f"{MONTHS_RU[key[1] - 1]} {key[0]}")
            month_cell.font = Font(bold=True, color="FFFFFF", size=12)
            month_cell.fill = PatternFill("solid", fgColor="5B9BD5")
            month_cell.alignment = Alignment(horizontal="left")
            row_number += 1

        stage = stages[lead["status_id"]]
        row_values = lead_row(
            lead, stage, users, config["custom_fields"], contacts, companies, notes
        )
        for column, value in enumerate(row_values, 1):
            cell = ws.cell(row_number, column, value)
            cell.fill = PatternFill("solid", fgColor=stage.color)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
        ws.cell(row_number, 8).number_format = '#,##0.00 "BYN"'
        row_number += 1

        next_month = month_key(valid[index + 1]) if index + 1 < len(valid) else None
        if next_month != key:
            totals = month_totals[key]
            ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=2)
            ws.cell(row_number, 1, "Сумма в работе").font = Font(bold=True, color="1F4E78")
            ws.cell(row_number, 3, totals["work"]).number_format = '#,##0.00 "BYN"'
            ws.merge_cells(start_row=row_number, start_column=4, end_row=row_number, end_column=5)
            ws.cell(row_number, 4, "Сумма успешных").font = Font(bold=True, color="006100")
            ws.cell(row_number, 6, totals["won"]).number_format = '#,##0.00 "BYN"'
            for column in range(1, len(HEADERS) + 1):
                ws.cell(row_number, column).fill = PatternFill("solid", fgColor="EAF2F8")
                ws.cell(row_number, column).border = Border(bottom=thin)
            row_number += 1

    widths = [32, 24, 25, 28, 24, 18, 25, 16, 24, 48, 35]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 42
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def selected_pipelines(all_pipelines: Iterable[dict], selection: list[Any]) -> list[dict]:
    if not selection:
        return list(all_pipelines)
    ids = {int(x) for x in selection if str(x).isdigit()}
    names = {str(x).casefold() for x in selection if not str(x).isdigit()}
    result = [p for p in all_pipelines if p["id"] in ids or p.get("name", "").casefold() in names]
    missing = len(selection) - len(result)
    if missing > 0:
        print("Предупреждение: некоторые указанные воронки не найдены")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Excel-отчёты по сделкам amoCRM")
    parser.add_argument("--config", default="config.json")
    parser.add_argument("--env", default=".env")
    args = parser.parse_args()

    config_path = Path(args.config).resolve()
    env_path = Path(args.env).resolve()
    config = json.loads(config_path.read_text("utf-8"))
    client = AmoCRMClient(env_path)
    users = client.users()
    pipelines = selected_pipelines(client.pipelines(), config.get("pipelines", []))
    if not pipelines:
        raise RuntimeError("Не найдено ни одной воронки для отчёта")

    date_from = parse_date(config["date_from"])
    date_to = parse_date(config.get("date_to"), end_of_day=True)
    output_dir = (config_path.parent / config.get("output_dir", "reports")).resolve()

    for pipeline in pipelines:
        stages = pipeline_stages(pipeline, config)
        leads = client.leads(pipeline["id"], date_from, date_to)
        contact_ids = {
            item_id for lead in leads for item_id in related_ids(lead, "contacts")
        }
        company_ids = {
            item_id for lead in leads for item_id in related_ids(lead, "companies")
        }
        contacts = client.contacts(contact_ids)
        companies = client.companies(company_ids)
        notes = client.lead_notes({lead["id"] for lead in leads})
        filename = f"{sanitize_filename(pipeline['name'])}.xlsx"
        create_workbook(
            pipeline["name"], leads, stages, users,
            contacts, companies, notes, config, output_dir / filename,
        )
        print(f"Создан: {output_dir / filename} — {len(leads)} сделок получено")


if __name__ == "__main__":
    main()
